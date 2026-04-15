#!/usr/bin/env python3
"""Plot episodic rewards from tempresults.xlsx for p=-1 and p=0.5, seed=1."""
import ast
import os
import sys

import numpy as np

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    import matplotlib.pyplot as plt
except ImportError:
    print("Install matplotlib: pip install matplotlib", file=sys.stderr)
    sys.exit(1)


def parse_returns(s):
    if s is None or (isinstance(s, str) and s.strip() == ""):
        return []
    if isinstance(s, (list, tuple)):
        return [float(x) for x in s]
    try:
        out = ast.literal_eval(s)
        return [float(x) for x in out] if out else []
    except (ValueError, SyntaxError, TypeError):
        pass
    # Excel truncates at 32767 chars; parse as comma-separated floats
    text = s.strip().lstrip("[").rstrip("]")
    vals = []
    for tok in text.split(","):
        tok = tok.strip()
        if not tok:
            continue
        try:
            vals.append(float(tok))
        except ValueError:
            break
    return vals


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(script_dir, "tempresults.xlsx")
    if not os.path.isfile(path):
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        print("Empty spreadsheet.")
        sys.exit(0)

    headers = [str(h).strip() if h else f"_col{j}" for j, h in enumerate(rows[0])]
    col = {h: i for i, h in enumerate(headers)}

    target_seed = 1
    target_probs = {-1.0, 0.5}

    fig, ax = plt.subplots(figsize=(12, 6))
    colors = {-1.0: "tab:blue", 0.5: "tab:orange"}
    labels_seen = set()

    for row in rows[1:]:
        if row[0] is None:
            break
        seed = row[col["seed"]]
        prob = row[col["replay_resample_prob"]]
        if seed != target_seed or prob not in target_probs:
            continue

        returns = parse_returns(row[col["episodic_returns"]])
        if not returns:
            continue

        run_name = row[col["run_name"]]
        episodes = np.arange(1, len(returns) + 1)
        group_label = f"p = {prob}"
        label = group_label if prob not in labels_seen else None
        labels_seen.add(prob)

        ax.plot(episodes, returns, color=colors[prob], alpha=0.7, linewidth=1,
                label=label)

    ax.set_xlabel("Episode")
    ax.set_ylabel("Episodic Return")
    ax.set_title("Episodic Returns: p = -1 vs p = 0.5 (seed = 1)")
    ax.legend()
    ax.grid(True, alpha=0.3)

    plt.tight_layout()
    out_path = os.path.join(script_dir, "tempresults_graph.png")
    plt.savefig(out_path, dpi=150, bbox_inches="tight")
    print(f"Saved: {out_path}")
    plt.show()


if __name__ == "__main__":
    main()
